#!/usr/bin/env python3
"""Build the one-slide "Token Utilisation Plan" PPTX from a normalised JSON input.

Usage:
    python build_report.py input.json output.pptx [--template path/to/template.pptx]

The business rules come from "Token Utilisation Plan | Master Prompt v3":
  * commitment is MONTHLY; "Tokens Remaining" is never calculated or shown
  * status = Actual vs Planned for the current reporting month (not vs commitment)
  * utilisation = Actual / Monthly commitment
  * future months are planned only; consumption is never cumulative
  * nothing is invented: optional blocks that have no data are left out
"""
import argparse
import copy
import json
import sys
from datetime import date
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = HERE.parent / "assets" / "template.pptx"

# Palette: Genesys-style orange/coral primary, deep navy text, restrained accents.
NAVY = RGBColor(0x0E, 0x28, 0x41)
CORAL = RGBColor(0xE9, 0x71, 0x32)
CORAL_LIGHT = RGBColor(0xF9, 0xC9, 0xB0)
GREY = RGBColor(0x6B, 0x7A, 0x8F)
GREY_BAR = RGBColor(0x9A, 0xA5, 0xB8)
CARD_BG = RGBColor(0xF6, 0xF8, 0xFB)
CARD_LINE = RGBColor(0xDD, 0xE3, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PURPLE = RGBColor(0x6A, 0x5A, 0xCD)
BLUE = RGBColor(0x0F, 0x9E, 0xD5)
TEAL = RGBColor(0x14, 0xB8, 0xA6)
HILITE = RGBColor(0xFF, 0xF1, 0xB8)

STATUS = {
    "GREEN": ("On Track", RGBColor(0xE6, 0xF4, 0xEA), RGBColor(0x2E, 0x9E, 0x57)),
    "YELLOW": ("Slightly Behind Plan", RGBColor(0xFF, 0xF4, 0xD6), RGBColor(0xF2, 0xA9, 0x00)),
    "RED": ("Behind Plan", RGBColor(0xFD, 0xE8, 0xE4), RGBColor(0xD9, 0x30, 0x25)),
}
STEPS = [
    ("1. ACTIVATE", "Launch priority use cases and drive initial adoption", CORAL, RGBColor(0xFD, 0xEB, 0xE2)),
    ("2. SCALE", "Increase usage across teams and channels", PURPLE, RGBColor(0xEC, 0xEA, 0xFB)),
    ("3. EXPAND", "Add workloads, teams and products", BLUE, RGBColor(0xE1, 0xF3, 0xFC)),
    ("4. OPTIMIZE", "Sustain utilisation and maximize value", TEAL, RGBColor(0xDD, 0xF5, 0xF1)),
]
LANE_COLORS = [(CORAL, RGBColor(0xFB, 0xD9, 0xC8)), (PURPLE, RGBColor(0xDD, 0xD8, 0xF7)),
               (TEAL, RGBColor(0xC8, 0xEF, 0xEA)), (BLUE, RGBColor(0xCB, 0xE8, 0xF7))]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


# --------------------------------------------------------------------------- data
def ym(s):
    y, m = str(s).split("-")[:2]
    return int(y), int(m)


def ym_idx(t):
    return t[0] * 12 + t[1] - 1


def idx_ym(i):
    return i // 12, i % 12 + 1


def ym_key(t):
    return f"{t[0]:04d}-{t[1]:02d}"


def month_label(t, short_year=True):
    return f"{MONTHS[t[1] - 1]} {str(t[0])[2:] if short_year else t[0]}"


def load(path):
    d = json.loads(Path(path).read_text(encoding="utf-8"))
    missing = [k for k in ("customer", "monthly_commitment", "contract_end", "reporting_month")
               if d.get(k) in (None, "")]
    if missing:
        sys.exit(f"ERROR: missing required fields: {', '.join(missing)}")
    d["actual"] = {ym_key(ym(k)): float(v) for k, v in (d.get("actual") or {}).items()}
    d["planned"] = {ym_key(ym(k)): float(v) for k, v in (d.get("planned") or {}).items()}
    rep = ym_key(ym(d["reporting_month"]))
    d["current_actual"] = float(d["current_actual"]) if d.get("current_actual") is not None else d["actual"].get(rep)
    d["current_planned"] = float(d["current_planned"]) if d.get("current_planned") is not None else d["planned"].get(rep)
    if d["current_actual"] is None or d["current_planned"] is None:
        sys.exit("ERROR: actual and planned consumption for the current reporting month are required "
                 "(provide them in 'actual'/'planned' or as current_actual/current_planned).")
    if float(d["current_planned"]) == 0 or float(d["monthly_commitment"]) == 0:
        sys.exit("ERROR: planned consumption and monthly commitment must be non-zero.")
    d["monthly_commitment"] = float(d["monthly_commitment"])
    return d


def metrics(d):
    rep = ym(d["reporting_month"])
    end = ym(d["contract_end"][:7])
    a, p, c = d["current_actual"], d["current_planned"], d["monthly_commitment"]
    var = a - p
    var_pct = var / p
    th = d.get("thresholds", {"green": -0.05, "yellow": -0.10})
    code = "GREEN" if var_pct >= th["green"] else "YELLOW" if var_pct >= th["yellow"] else "RED"
    months_left = ym_idx(end) - ym_idx(rep)
    return dict(rep=rep, end=end, actual=a, planned=p, commit=c, var=var, var_pct=var_pct,
                util=a / c, status=code, months_left=months_left,
                first_left=idx_ym(ym_idx(rep) + 1))


def unit_for(commit):
    if commit >= 1e6:
        return 1e6, "M"
    if commit >= 1e3:
        return 1e3, "K"
    return 1.0, ""


def fmt(v, unit, signed=False, dec=None):
    div, suf = unit
    x = v / div
    if dec is None:
        dec = 1 if suf else 0
        if suf and abs(x) < 10 and abs(x - round(x, 1)) > 1e-9:
            dec = 2
    s = f"{x:+,.{dec}f}" if signed else f"{x:,.{dec}f}"
    return s + suf


# --------------------------------------------------------------------------- drawing
def box(slide, x, y, w, h, fill=CARD_BG, line=CARD_LINE, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: list of (text, size, bold, color) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, (t, size, bold, color) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def shape_text(shape, t, size, bold, color, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = t
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


def header(slide, d, m, u):
    text(slide, 0.4, 0.15, 7.9, 0.5, [(d["customer"] + " — Token Utilisation Plan", 22, True, NAVY)])
    # colour the title suffix like the reference layout
    tb = slide.shapes[-1]
    p = tb.text_frame.paragraphs[0]
    p.runs[0].text = d["customer"] + " "
    r = p.add_run()
    r.text = "— Token Utilisation Plan"
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = CORAL
    text(slide, 0.4, 0.74, 7.9, 0.3,
         [("A clear plan to reach and sustain monthly token utilisation through contract expiration.", 11, False, GREY)])

    label, bg, fg = STATUS[m["status"]]
    box(slide, 8.3, 0.12, 3.0, 0.95, fill=bg, line=fg)
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.45), Inches(0.38), Inches(0.42), Inches(0.42))
    dot.fill.solid()
    dot.fill.fore_color.rgb = fg
    dot.line.fill.background()
    text(slide, 8.97, 0.17, 2.3, 0.4, [(label, 15, True, NAVY)])
    text(slide, 8.97, 0.55, 2.3, 0.5, [
        (f"{fmt(m['var'], u, signed=True)} tokens | {m['var_pct'] * 100:+.1f}% vs plan", 9.5, True, fg),
        (f"{month_label(m['rep'], False)} actual vs planned", 8, False, GREY)])

    end_dt = date.fromisoformat(d["contract_end"][:10]) if len(d["contract_end"]) >= 10 else None
    end_txt = end_dt.strftime("%b %d, %Y") if end_dt else month_label(m["end"], False)
    text(slide, 11.5, 0.14, 1.5, 0.95, [
        ("Contract End Date", 8, False, GREY), (end_txt, 12, True, NAVY),
        ("Current Period", 8, False, GREY), (month_label(m["rep"], False), 11, True, NAVY)])


def kpis(slide, d, m, u):
    rep = month_label(m["rep"], False)
    cards = [
        ("Monthly Token Commitment", fmt(m["commit"], u), "tokens per month", NAVY),
        (f"{rep} Planned Consumption", fmt(m["planned"], u), "tokens", NAVY),
        (f"{rep} Actual Consumption", fmt(m["actual"], u), "tokens", NAVY),
        (f"Variance to Plan ({rep})", fmt(m["var"], u, signed=True),
         f"tokens ({m['var_pct'] * 100:+.1f}%)", STATUS[m["status"]][2] if m["var"] < 0 else NAVY),
        (f"Current Monthly Utilisation ({rep})", f"{m['util'] * 100:.0f}%",
         f"{fmt(m['actual'], u)} / {fmt(m['commit'], u)} tokens", NAVY),
        ("Months Remaining", str(m["months_left"]),
         f"{month_label(m['first_left'], False)} – {month_label(m['end'], False)}" if m["months_left"] > 0 else "Final month", NAVY),
    ]
    gap, x0, y, h = 0.12, 0.35, 1.2, 1.05
    w = (13.333 - 2 * x0 - 5 * gap) / 6
    for i, (lab, val, sub, col) in enumerate(cards):
        x = x0 + i * (w + gap)
        box(slide, x, y, w, h)
        text(slide, x + 0.12, y + 0.07, w - 0.24, 0.3, [(lab, 8.5, True, NAVY)])
        text(slide, x + 0.12, y + 0.4, w - 0.24, 0.35, [(val, 20, True, col)])
        text(slide, x + 0.12, y + 0.77, w - 0.24, 0.2, [(sub, 8, False, GREY)])
        if i == 4:  # utilisation progress bar (against the monthly commitment)
            box(slide, x + 0.12, y + 1.0 - 0.02, w - 0.24, 0.035, fill=CARD_LINE, line=None, radius=0.5)
            box(slide, x + 0.12, y + 1.0 - 0.02, (w - 0.24) * min(m["util"], 1), 0.035, fill=CORAL, line=None, radius=0.5)


def path_band(slide):
    y, h = 2.37, 0.78
    box(slide, 0.35, y, 12.63, h, fill=WHITE)
    text(slide, 0.5, y + 0.1, 1.45, 0.6, [("Our Path to Sustained Utilisation", 11, True, NAVY)], anchor=MSO_ANCHOR.MIDDLE)
    x, w = 2.05, 2.72
    for i, (title, desc, col, bg) in enumerate(STEPS):
        s = box(slide, x + i * (w - 0.12), y + 0.08, w, h - 0.16, fill=bg, line=None, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
        s.adjustments[0] = 0.25
        text(slide, x + i * (w - 0.12) + (0.45 if i else 0.25), y + 0.13, w - 0.95, 0.55,
             [(title, 10.5, True, col), (desc, 8.5, False, NAVY)], anchor=MSO_ANCHOR.MIDDLE)


def _set_line(ser_el, rgb, dash=None, width_pt=1.75):
    ln = etree.SubElement(etree.SubElement(ser_el, qn("c:spPr")), qn("a:ln"), w=str(int(width_pt * 12700)))
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=str(rgb))
    if dash:
        etree.SubElement(ln, qn("a:prstDash"), val=dash)
    return ln


def chart(slide, d, m, u):
    x, y, w, h = 0.35, 3.25, 12.63, 2.4
    box(slide, x, y, w, h, fill=WHITE)
    text(slide, x + 0.15, y + 0.07, 5.5, 0.3,
         [(f"Monthly Token Consumption vs Plan (Tokens{' in ' + {'M': 'Millions', 'K': 'Thousands'}[u[1]] if u[1] else ''})", 12, True, NAVY)])

    first = min([ym_idx(ym(k)) for k in list(d["actual"]) + list(d["planned"])] + [ym_idx(m["rep"])])
    last = ym_idx(m["end"])
    months = [idx_ym(i) for i in range(first, last + 1)]
    rep_i = months.index(m["rep"])
    div = u[0]

    cd = CategoryChartData(number_format="0.0")
    cd.categories = [month_label(t) for t in months]
    actual = [d["actual"].get(ym_key(t)) / div if (ym_key(t) in d["actual"] and i <= rep_i) else None
              for i, t in enumerate(months)]
    planned = [d["planned"].get(ym_key(t)) / div if ym_key(t) in d["planned"] else None for t in months]
    cd.add_series("Actual Consumption", actual)
    cd.add_series("Planned Consumption", planned)
    cd.add_series(f"Monthly Commitment ({fmt(m['commit'], u)})", [m["commit"] / div] * len(months))

    cx, cy, cw, ch = x + 0.1, y + 0.35, w - 0.2, h - 0.4
    inner = (0.05, 0.23, 0.94, 0.60)  # x, y, w, h fractions of the chart frame
    # highlight band for the current month, drawn behind the chart
    band = cw * inner[2] / len(months)
    hl = box(slide, cx + cw * inner[0] + rep_i * band, cy + ch * inner[1] - 0.12, band, ch * inner[3] + 0.12,
             fill=HILITE, line=None, radius=0.05)
    pill = box(slide, cx + cw * inner[0] + rep_i * band - 0.25, cy + ch * inner[1] - 0.2, band + 0.5, 0.2,
               fill=RGBColor(0xF5, 0xC3, 0x42), line=None, radius=0.4)
    shape_text(pill, "Current Month", 7, True, NAVY)

    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(cx), Inches(cy), Inches(cw), Inches(ch), cd)
    ch_ = gf.chart
    ch_.font.size = Pt(8)
    ch_.font.color.rgb = GREY
    ch_.has_legend = True
    ch_.legend.position = XL_LEGEND_POSITION.TOP
    ch_.legend.include_in_layout = False
    ch_.legend.font.size = Pt(8)
    plot = ch_.plots[0]
    plot.gap_width = 40
    plot.overlap = 0
    for ser, col in zip(plot.series, (GREY_BAR, CORAL_LIGHT)):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
    va = ch_.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = CARD_LINE
    va.format.line.fill.background()
    va.minimum_scale = 0
    va.maximum_scale = max([m["commit"] / div] + [v for v in actual + planned if v] + [0]) * 1.2
    va.tick_labels.font.size = Pt(8)
    ch_.category_axis.tick_labels.font.size = Pt(7.5)
    ch_.category_axis.format.line.color.rgb = CARD_LINE

    # data labels: actual for history, planned for months without an actual bar
    for ser, keep in ((plot.series[0], [v is not None for v in actual]),
                      (plot.series[1], [pv is not None and av is None for pv, av in zip(planned, actual)])):
        dl = ser.data_labels
        dl.show_value = True
        dl.number_format = "0.0" if u[1] else "#,##0"
        dl.number_format_is_linked = False
        dl.font.size = Pt(7)
        dl.font.color.rgb = NAVY
        dls = ser._element.find(qn("c:dLbls"))
        for i in reversed(range(len(months))):
            if not keep[i]:
                e = etree.Element(qn("c:dLbl"))
                etree.SubElement(e, qn("c:idx"), val=str(i))
                etree.SubElement(e, qn("c:delete"), val="1")
                dls.insert(0, e)

    # move the commitment series into a line chart that shares the same axes (combo chart)
    cs = ch_._chartSpace
    plot_area = cs.find(".//" + qn("c:plotArea"))
    bar = plot_area.find(qn("c:barChart"))
    sers = bar.findall(qn("c:ser"))
    commit = sers[2]
    bar.remove(commit)
    for tag in ("c:spPr", "c:invertIfNegative", "c:dLbls"):
        for el in commit.findall(qn(tag)):
            commit.remove(el)
    sp = etree.Element(qn("c:spPr"))
    ln = etree.SubElement(sp, qn("a:ln"), w="22225")
    etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr"), val=str(CORAL))
    etree.SubElement(ln, qn("a:prstDash"), val="dash")
    commit.insert(3, sp)  # after idx, order, tx
    mk = etree.Element(qn("c:marker"))
    etree.SubElement(mk, qn("c:symbol"), val="none")
    commit.insert(4, mk)
    line = etree.Element(qn("c:lineChart"))
    etree.SubElement(line, qn("c:grouping"), val="standard")
    etree.SubElement(line, qn("c:varyColors"), val="0")
    line.append(commit)
    etree.SubElement(line, qn("c:marker"), val="1")
    for ax in bar.findall(qn("c:axId")):
        etree.SubElement(line, qn("c:axId"), val=ax.get("val"))
    bar.addnext(line)

    # fixed inner plot area so the highlight band lines up with the categories
    lay = plot_area.find(qn("c:layout"))
    if lay is None:
        lay = etree.Element(qn("c:layout"))
        plot_area.insert(0, lay)
    for c in list(lay):
        lay.remove(c)
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    etree.SubElement(ml, qn("c:layoutTarget"), val="inner")
    for tag, v in (("xMode", "edge"), ("yMode", "edge"), ("x", inner[0]), ("y", inner[1]), ("w", inner[2]), ("h", inner[3])):
        etree.SubElement(ml, qn(f"c:{tag}"), val=str(v))


def use_cases(slide, d, m, x, y, w, h):
    ucs = (d.get("use_cases") or [])[:4]
    if len(d.get("use_cases") or []) > 4:
        print("WARNING: only the first 4 use cases are drawn.", file=sys.stderr)
    box(slide, x, y, w, h, fill=WHITE)
    text(slide, x + 0.15, y + 0.06, 4, 0.25, [("Use Case Activation Plan", 12, True, NAVY)])
    q0 = (m["rep"][0], (m["rep"][1] - 1) // 3 * 3 + 1)
    qe = (m["end"][0], (m["end"][1] - 1) // 3 * 3 + 1)
    quarters = []
    i = ym_idx(q0)
    while i <= ym_idx(qe):
        quarters.append(idx_ym(i))
        i += 3
    left, top = x + 1.75, y + 0.34
    tw = w - 1.75 - 0.1
    qw = tw / len(quarters)
    for k, q in enumerate(quarters):
        text(slide, left + k * qw, top, qw, 0.28,
             [(f"Q{(q[1] - 1) // 3 + 1} {q[0]}", 8, True, NAVY)], align=PP_ALIGN.CENTER)
    lane_h = min(0.4, (h - 0.66) / max(len(ucs), 1))
    for r, uc in enumerate(ucs):
        ly = top + 0.26 + r * lane_h
        dark, light = LANE_COLORS[r % len(LANE_COLORS)]
        text(slide, x + 0.15, ly, 1.55, lane_h, [(uc["name"], 9.5, True, NAVY)]
             + ([(uc["description"], 7, False, GREY)] if uc.get("description") and lane_h >= 0.38 else []), anchor=MSO_ANCHOR.MIDDLE)
        for ms in uc.get("milestones", []):
            s_i, e_i = ym_idx(ym(ms["start"])), ym_idx(ym(ms.get("end") or ms["start"]))
            t0, t1 = ym_idx(quarters[0]), ym_idx(quarters[-1]) + 3
            s_i, e_i = max(s_i, t0), min(e_i + 1, t1)
            if e_i <= s_i:
                continue
            bx = left + (s_i - t0) / (t1 - t0) * tw
            bw = (e_i - s_i) / (t1 - t0) * tw
            c = box(slide, bx, ly + 0.04, bw - 0.03, lane_h - 0.08, fill=light, line=None, shape=MSO_SHAPE.PENTAGON)
            c.adjustments[0] = 0.3
            shape_text(c, ms["label"], 8, False, NAVY)


def focus(slide, d, x, y, w, h, columns=False):
    items = (d.get("focus_areas") or [])[:4]
    box(slide, x, y, w, h, fill=WHITE)
    text(slide, x + 0.15, y + 0.06, 3, 0.25, [("Key Focus Areas", 12, True, NAVY)])
    colors = [CORAL, PURPLE, BLUE, TEAL]
    if columns:
        cw = (w - 0.3) / max(len(items), 1)
        for i, it in enumerate(items):
            text(slide, x + 0.15 + i * cw, y + 0.4, cw - 0.15, h - 0.45,
                 [(it["title"], 10, True, colors[i % 4]), (it["text"], 8.5, False, NAVY)])
        return
    rh = (h - 0.4) / max(len(items), 1)
    for i, it in enumerate(items):
        ry = y + 0.38 + i * rh
        bar = box(slide, x + 0.15, ry + 0.03, 0.05, rh - 0.1, fill=colors[i % 4], line=None, radius=0.5)
        big = len(items) <= 2
        text(slide, x + 0.3, ry, w - 0.45, rh,
             [(it["title"], 9.5 if big else 8.5, True, NAVY), (it["text"], 8 if big else 7, False, GREY)], anchor=MSO_ANCHOR.MIDDLE)


def build(d, out, template):
    m = metrics(d)
    u = unit_for(m["commit"])
    prs = Presentation(str(template))
    # drop the template's sample slide, keep its theme/size, add a blank slide
    sld = prs.slides._sldIdLst
    for s in list(sld):
        prs.part.drop_rel(s.rId)
        sld.remove(s)
    blank = next((l for l in prs.slide_layouts if l.name.lower() == "blank"), prs.slide_layouts[-1])
    slide = prs.slides.add_slide(blank)
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)

    header(slide, d, m, u)
    kpis(slide, d, m, u)
    path_band(slide)
    chart(slide, d, m, u)
    by, bh = 5.75, 1.47
    has_uc, has_fa = bool(d.get("use_cases")), bool(d.get("focus_areas"))
    if has_uc and has_fa:
        use_cases(slide, d, m, 0.35, by, 8.4, bh)
        focus(slide, d, 8.85, by, 4.13, bh)
    elif has_uc:
        use_cases(slide, d, m, 0.35, by, 12.63, bh)
    elif has_fa:
        focus(slide, d, 0.35, by, 12.63, bh, columns=True)
    src = d.get("source_note") or "Source: Tableau exports. Tokens are committed monthly; unused tokens do not carry forward."
    text(slide, 0.35, 7.25, 12.6, 0.2, [(src, 7, False, GREY)])
    prs.save(out)
    return m, u


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output")
    ap.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    a = ap.parse_args()
    d = load(a.input)
    m, u = build(d, a.output, a.template)
    print(f"Saved {a.output}")
    print(f"Customer: {d['customer']} | period {month_label(m['rep'], False)} | contract end {month_label(m['end'], False)}")
    print(f"Actual {fmt(m['actual'], u)} vs planned {fmt(m['planned'], u)} -> variance "
          f"{fmt(m['var'], u, signed=True)} ({m['var_pct'] * 100:+.1f}%) -> {m['status']}")
    print(f"Utilisation vs monthly commitment {fmt(m['commit'], u)}: {m['util'] * 100:.1f}% | months remaining: {m['months_left']}")


if __name__ == "__main__":
    main()
